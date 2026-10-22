#!/usr/bin/env python3
"""Rebuild evidence from shipped sources. Required missing steps never become PASS."""
from pathlib import Path
from dataclasses import asdict
import sys,os,json,csv,hashlib,subprocess,shutil,tempfile,xml.etree.ElementTree as ET,argparse
from datetime import date,datetime,timezone
BASE=Path(__file__).resolve().parent;ROOT=BASE.parent
sys.path.insert(0,str(BASE/'python'))
if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')

def execute(cmd,log,cwd=None,timeout=180):
    p=subprocess.run([str(c) for c in cmd],cwd=cwd,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=timeout)
    log.write_text('COMMAND: '+subprocess.list2cmdline([str(c) for c in cmd])+'\nEXIT_CODE: '+str(p.returncode)+'\n'+p.stdout+'\n'+p.stderr,encoding='utf-8')
    if p.returncode:raise RuntimeError(f'Command failed ({p.returncode}); see {log.name}')
    return p.stdout

def validate_spec(path):
    from openpyxl import load_workbook
    from atpm_gostay.ch1_kripke.gostay_model import build_model,verify_gostay_kripke
    wb=load_workbook(path,read_only=True,data_only=True)
    expected={'ThongTin','YeuCauChucNang','YeuCauPhiChucNang','RangBuoc','MoHinhTrangThai','ThuocTinhAnToan','LoaiTruLanNhau','HamSinhMa','KiemTraSoHoc','TestCase'}
    if set(wb.sheetnames)!=expected:raise ValueError('Workbook must use the documented ten-sheet schema')
    def rows(name):
        data=list(wb[name].values);headers=data[0]
        return [dict(zip(headers,row)) for row in data[1:] if any(x is not None for x in row)]
    reqs={r['ID'] for name in ('YeuCauChucNang','YeuCauPhiChucNang','RangBuoc') for r in rows(name)}
    for tc in rows('TestCase'):
        for key in tc['YeuCauLienQuan'].split(','):
            if key.strip() not in reqs:raise ValueError('Dangling requirement '+key)
    properties=[r['CongThucCTL'] for r in rows('ThuocTinhAnToan')]
    for version in ('v0','v1'):
        model=build_model(version);prefix=version+'_'
        states=[r for r in rows('MoHinhTrangThai') if r['TrangThai'].startswith(prefix)]
        if {r['TrangThai'][len(prefix):] for r in states}!=model.states:raise ValueError('State set differs from code')
        for r in states:
            state=r['TrangThai'][len(prefix):]
            labels={s.strip() for s in r['Nhan'].split(',')}
            targets={s.strip()[len(prefix):] for s in r['ChuyenDenCacTrangThai'].split(',')}
            if labels!=set(model.labels_of(state)) or targets!=model.transitions[state]:raise ValueError('Transition/label mismatch: '+state)
        initial={r['TrangThai'][len(prefix):] for r in states if r.get('KhoiTao')=='x'}
        if initial!=model.initial_states:raise ValueError('Initial states differ')
        if set(properties)!=set(verify_gostay_kripke(version)['properties']):raise ValueError('CTL list differs')
    if len(rows('KiemTraSoHoc'))!=5:raise ValueError('Expected five documented SMT checks')
    return {'status':'PASS','sheets':10,'requirements':len(reqs),'test_cases':len(rows('TestCase')),
            'note':'Workbook is validated by this runner; HamSinhMa documents existing C examples, not automatic generation.'}

def main():
    args=argparse.ArgumentParser();args.add_argument('--strict-native',action='store_true')
    args.add_argument('--output',type=Path,help='Save this run separately; defaults to 04_Ket_qua')
    options=args.parse_args()
    from atpm_gostay.ch1_kripke.gostay_model import verify_gostay_kripke,build_model
    from atpm_gostay.ch3_smt_bmc.z3_bmc_verifier import check_inventory_conservation,check_int32_overflow,check_date_inversion,check_quantity_limit,check_range_limit,check_bmc
    from atpm_gostay.ch4_fuzzing.dynamic_fuzzer import write_corpus,summarize_java_results
    from atpm_gostay.ch2_memory_safety.memory_checker import run_c_memory_lab
    out=options.output.resolve() if options.output else ROOT/'04_Ket_qua';data=out/'data';logs=out/'logs'
    data.mkdir(parents=True,exist_ok=True);logs.mkdir(parents=True,exist_ok=True)
    report={'generated_utc':datetime.now(timezone.utc).isoformat(),'status':'RUNNING','failures':[],'limits':[
        'Java corpus runs on actual services with mocked repositories, not live HTTP.',
        'Database concurrency regression uses Hibernate/H2; PostgreSQL concurrency is not asserted.',
        'CTL liveness assumes successful eventual scheduler execution.',
        'BMC is bounded: one order, quantity 1, capacity 10. No completeness threshold is claimed.',
        'C models deliberately contain pointer bugs; these are not Java heap vulnerabilities.']}
    import platform,z3
    report['environment']={'python':platform.python_version(),'z3':z3.get_version_string(),
        'platform':platform.platform(),'java':subprocess.run(['java','-version'],capture_output=True,text=True).stderr.splitlines()[0]}
    def step(name,fn):
        print('RUN '+name,flush=True)
        try:report[name]=fn()
        except Exception as exc:report[name]={'status':'ERROR','message':str(exc)};report['failures'].append(name)
    step('spec',lambda:validate_spec(ROOT/'01_Bao_cao/spec_gostay_inventory_lock_da_sua.xlsx'))
    step('kripke',lambda:[verify_gostay_kripke(v) for v in ('v0','v1')])
    step('smt',lambda:[asdict(f(v)) for v in ('v0','v1') for f in (check_inventory_conservation,check_int32_overflow,check_date_inversion,check_quantity_limit,check_range_limit)])
    step('bmc',lambda:[check_bmc(v,k) for k in (0,1,2,3,6) for v in ('v0','v1')])
    for version in ('v0','v1'):(data/f'kripke_{version}.dot').write_text(build_model(version).to_dot(),encoding='utf-8')
    write_corpus(data/'shared_corpus.csv')
    report['corpus_sha256']=hashlib.sha256((data/'shared_corpus.csv').read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix='gostay-native-') as native:
        step('memory',lambda:run_c_memory_lab(BASE/'c/ch2_memory',Path(native)))
        for f in Path(native).glob('*.log'):shutil.copy2(f,logs/f.name)
    def unit_tests():
        target=data/'python_tests.xml'
        execute([sys.executable,'-m','pytest',str(BASE/'python/tests'),'-q',f'--junitxml={target}'],logs/'python_tests.log',BASE)
        root=ET.parse(target).getroot();suites=list(root.iter('testsuite'))
        return {'status':'PASS','tests':sum(int(s.get('tests',0)) for s in suites),'failures':sum(int(s.get('failures',0))+int(s.get('errors',0)) for s in suites)}
    step('python',unit_tests)
    def ast_analysis():
        javac=shutil.which('javac');java=shutil.which('java')
        if not javac or not java:raise RuntimeError('JDK javac/java missing')
        counts={}
        with tempfile.TemporaryDirectory(prefix='gostay-ast-') as tmp:
            execute([javac,'-encoding','UTF-8','-d',tmp,BASE/'java/JavaAstAnalyzer.java'],logs/'ast_compile.log')
            for v in ('v0','v1'):
                source=ROOT/f'02_Source_Java/{v}/BookingandInventory/src/main/java'
                execute([java,'-cp',tmp,'JavaAstAnalyzer',v,source,data/f'ast_{v}.csv'],logs/f'ast_{v}.log')
                counts[v]=len(list(csv.DictReader((data/f'ast_{v}.csv').open(encoding='utf-8'))))
        return {'status':'PASS','methods':counts,'metric':'CC = 1 + if/loops/catch/ternary/short-circuit AND/OR; depth counts nested control constructs.'}
    step('ast',ast_analysis)
    def java_tests():
        results={};anchor=date.today().isoformat()
        for v in ('v0','v1'):
            source=ROOT/f'02_Source_Java/{v}/BookingandInventory'
            with tempfile.TemporaryDirectory(prefix='gostay-java-') as tmp:
                project=Path(tmp)/'project'
                shutil.copytree(source,project,ignore=shutil.ignore_patterns('target','.git'))
                dest=project/'src/test/java/com/Gostay/BookingandInventory/service/InventoryCorpusTest.java'
                shutil.copy2(BASE/'java/InventoryCorpusTest.java',dest)
                if v=='v0':
                    # Test instrumentation only. The shipped baseline application source is untouched.
                    shutil.copy2(BASE/'java/InventoryConcurrencyBaselineTest.java',dest.with_name('InventoryConcurrencyBaselineTest.java'))
                    pom=project/'pom.xml'
                    pom.write_text(pom.read_text(encoding='utf-8').replace('</dependencies>',
                        '<dependency><groupId>com.h2database</groupId><artifactId>h2</artifactId><scope>test</scope></dependency></dependencies>',1),encoding='utf-8')
                wrapper=project/('mvnw.cmd' if os.name=='nt' else 'mvnw')
                cmd=([os.environ.get('COMSPEC','cmd.exe'),'/d','/c',str(wrapper)] if os.name=='nt' else ['sh',str(wrapper)])
                cmd+=['-B','-Dstyle.color=never','-Dtest=Inventory*Test',f'-Datpm.version={v}',f'-Datpm.corpus={data/"shared_corpus.csv"}',
                      f'-Datpm.corpus.output={data/f"corpus_{v}_results.csv"}',f'-Datpm.anchor={anchor}','test']
                execute(cmd,logs/f'junit_{v}.log',project,timeout=240)
                suites=[]
                for f in sorted((project/'target/surefire-reports').glob('TEST-*.xml')):
                    x=ET.parse(f).getroot();suites.append({'name':x.get('name'),'tests':int(x.get('tests',0)),
                        'failures':int(x.get('failures',0)),'errors':int(x.get('errors',0))})
                    shutil.copy2(f,data/(v+'_'+f.name))
                results[v]={'status':'PASS','tests':sum(s['tests'] for s in suites),'suites':suites,
                    'corpus':summarize_java_results(data/f'corpus_{v}_results.csv')}
        return results
    step('java',java_tests)
    from atpm_gostay.ch1_kripke.draw_kripke import generate_figures
    step('figures',lambda:generate_figures(report,out/'hinh'))
    memory=report.get('memory',{})
    if any(x.get('status') in ('SKIPPED','ERROR') for x in memory.get('cppcheck',{}).values()):report['failures'].append('cppcheck_required')
    if memory.get('execution',{}).get('status')!='PASS':report['failures'].append('c_execution_required')
    if not memory.get('cppcheck',{}).get('v0',{}).get('findings'):report['failures'].append('c_v0_counterexample_missing')
    if memory.get('cppcheck',{}).get('v1',{}).get('findings'):report['failures'].append('c_v1_findings')
    if options.strict_native and memory.get('asan',{}).get('status')!='PASS':report['failures'].append('asan_required_by_flag')
    if any(r.get('status')=='UNKNOWN' for r in report.get('smt',[]) if isinstance(r,dict)):report['failures'].append('smt_unknown')
    if any(r['status']!=('SAT' if r['version']=='v0' else 'UNSAT') for r in report.get('smt',[]) if isinstance(r,dict)):
        report['failures'].append('smt_unexpected_result')
    if any(r['status']!=('SAT' if r['version']=='v0' and r['k']>=3 else 'UNSAT') for r in report.get('bmc',[]) if isinstance(r,dict)):
        report['failures'].append('bmc_unexpected_result')
    hashes={}
    for folder in ('02_Source_Java','03_Kiem_chung_ATPM'):
        for p in (ROOT/folder).rglob('*'):
            if p.is_file() and 'target' not in p.parts and '__pycache__' not in p.parts and p.suffix in ('.java','.py','.c','.h','.xml','.yaml','.json','.sh','.ps1'):
                hashes[str(p.relative_to(ROOT)).replace('\\','/')]=hashlib.sha256(p.read_bytes()).hexdigest()
    report['source_sha256']=hashes
    report['status']='FAIL' if report['failures'] else 'PASS_WITH_OPTIONAL_ASAN_SKIP' if memory.get('asan',{}).get('status')=='SKIPPED' else 'PASS'
    (data/'verification_results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    for key in ('kripke','smt','bmc'):(data/f'{key}_results.json').write_text(json.dumps(report.get(key),ensure_ascii=False,indent=2),encoding='utf-8')
    print('RESULT '+report['status']+'; details: '+str(data/'verification_results.json'))
    return 1 if report['failures'] else 0

if __name__=='__main__':sys.exit(main())
